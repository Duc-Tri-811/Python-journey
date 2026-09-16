#Day 13: Exercise 03: Create Report Sheet

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/sales.xlsx"

workbook = load_workbook(file_path)

report_sheet = workbook.create_sheet("Report")

report_sheet.append(["Sales Report"])
report_sheet.append(["'============"])
report_sheet.append(["Total Orders", 7])
report_sheet.append(["Total Revenue", 2420])
report_sheet.append(["Paid Orders", 4])
report_sheet.append(["Paid Revenue", 1620])

new_file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/sales_with_report.xlsx"
workbook.save(new_file_path)