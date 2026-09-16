#Day 13: Exercise 02: Add Revenue Category
from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

sheet["D1"] = "Revenue Category"

for index, row in enumerate(sheet.iter_rows(min_row=2,values_only=True), start=2):
    amount = row[1]

    if amount >= 300:
        category = "Large"
    else:
        category = "Small"

    sheet.cell(row=index, column=4).value = category

new_file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/sales_categorized.xlsx"
workbook.save(new_file_path)
        

