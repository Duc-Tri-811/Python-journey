#Day13: Exercise 05: Function

from openpyxl import load_workbook

def add_revenue_category(filename, output_filename):

    file_path = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/{filename}"

    workbook = load_workbook(file_path)
    sheet = workbook["Sales"]

    sheet["D1"] = "Revenue Category"

    for index, row in enumerate(sheet.iter_rows(min_row=2,values_only=True), start=2):
        amount = row[1]
        if amount >= 300:
            category = "Large"
        else:
            category = "Small"

        sheet.cell(row=index, column=4).value = category

    new_file_path = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/{output_filename}"
    workbook.save(new_file_path)

add_revenue_category("sales.xlsx","new_sales_categorized.xlsx")

     