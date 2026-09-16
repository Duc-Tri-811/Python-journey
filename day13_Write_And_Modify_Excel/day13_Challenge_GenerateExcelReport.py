#Day 13: Challenge: Generate_Excel_Report

def generate_excel_report(input_filename, output_filename):

    from openpyxl import load_workbook

    file_path = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/{input_filename}"

    workbook = load_workbook(file_path)
    sheet = workbook["Sales"]

    report_sheet = workbook.create_sheet("Report")

    total_orders = 0
    total_revenue = 0
    paid_orders = 0
    paid_revenue = 0
    largest_paid_order = 0
    for row in sheet.iter_rows(min_row=2,values_only=True):
        customer = row[0]
        amount = row[1]
        status = row[2]
        total_orders += 1
        total_revenue += amount
        if status == "paid":
            paid_orders += 1
            paid_revenue += amount
            if amount > largest_paid_order:
                largest_paid_order = amount
                largest_paid_customer = customer

    report_sheet.append(["Sales Report"])
    report_sheet.append(["Total Orders"])
    report_sheet.append(["Total Revenue"])
    report_sheet.append(["Paid Orders"])
    report_sheet.append(["Paid Revenue"])
    report_sheet.append(["Largest Paid Order"])
    report_sheet.append(["Largest Paid Customer"])
    
    report_sheet.cell(row=2, column=2).value = total_orders
    report_sheet.cell(row=3, column=2).value = total_revenue
    report_sheet.cell(row=4, column=2).value = paid_orders
    report_sheet.cell(row=5, column=2).value = paid_revenue
    report_sheet.cell(row=6, column=2).value = largest_paid_order
    report_sheet.cell(row=7, column=2).value = largest_paid_customer

    new_file_path = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/{output_filename}"
    workbook.save(new_file_path)

generate_excel_report("sales.xlsx","sales_automated_report.xlsx")

# Bonus Challenge
from openpyxl import load_workbook

new_file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day13_Write_And_Modify_Excel/sales_automated_report.xlsx"

workbook = load_workbook(new_file_path)
sheet = workbook["Sales"]

sheet["D1"] = "Revenue Category"

large_orders = 0
large_order_revenue = 0
for index, row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
    amount = row[1]
    if amount >= 300:
        category = "Large"
        large_orders += 1
        large_order_revenue += amount
    else:
        category = "Small"

    sheet.cell(row=index, column = 4).value = category

report_sheet = workbook["Report"]
report_sheet.append(["Large Orders"])
report_sheet.append(["Large Order Revenue"])

report_sheet.cell(row=8, column=2).value = large_orders
report_sheet.cell(row=9, column=2).value = large_order_revenue

workbook.save(new_file_path)



    




