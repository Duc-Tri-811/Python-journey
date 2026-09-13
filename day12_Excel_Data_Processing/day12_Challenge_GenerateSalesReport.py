#Day 12: Challenge: Generate sales report

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

def generate_sales_report(filename): 
    workbook = load_workbook(filename)
    sheet = workbook["Sales"]

    total_orders = 0
    total_revenue = 0
    paid_orders = 0
    paid_revenue = 0
    largest_paid_order = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):

#01: Calculat total orders
        total_orders += 1
#02: Calculate total revenue
        total_revenue += row[1]
#03: Calculate paid orders
        if row[2] == "paid":
            paid_orders += 1
#04: Calculate paid revenue
            paid_revenue += row[1]
#05: Find the largest paid order
            if largest_paid_order < row[1]:
                largest_paid_order = row[1]
                largest_paid_customer = row[0]


    return{
        "total_orders": total_orders,
        "total_revenue": total_revenue,
        "paid_orders": paid_orders,
        "paid_revenue": paid_revenue,
        "largest_paid_order": largest_paid_order,
        "largest_paid_customer": largest_paid_customer
    }

report = generate_sales_report(file_path)

print("Sales Report")
print("============")
print(f"Total orders: {report['total_orders']}")
print(f"Total revenue: {report['total_revenue']:.2f} €")
print(f"Paid orders: {report['paid_orders']}")
print(f"Paid revenue: {report['paid_revenue']:.2f} €")
print(f"Largest paid order: {report['largest_paid_customer']} - {report['largest_paid_order']} €")
