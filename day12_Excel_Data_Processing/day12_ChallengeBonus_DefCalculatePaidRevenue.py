#Day 12: Challenge Bonus: def Calculate paid revenue

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

def calculate_paid_revenue(filename):
    workbook = load_workbook(filename)
    sheet = workbook["Sales"]

    paid_revenue = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[2] == "paid":
            paid_revenue += row[1]

    return{
        "paid_revenue": paid_revenue
    }

result_paid_revenue = calculate_paid_revenue(file_path)

print(f"Paid revenue: {result_paid_revenue['paid_revenue']:.2f} €")