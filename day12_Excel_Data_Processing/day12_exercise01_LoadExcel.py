#Day 12: Exercise 01: Load Excel
from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

for row in sheet.iter_rows(values_only=True):
        print(row)