#Day 12: Exercise 04: Filter paid order
from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

paid_revenue = 0
for row in sheet.iter_rows(values_only=True):
    if row[2] == "paid":
        print(f"{row[0]} - {row[1]} €")
        paid_revenue += row[1]

print(f"Paid revenue: {paid_revenue} €")

