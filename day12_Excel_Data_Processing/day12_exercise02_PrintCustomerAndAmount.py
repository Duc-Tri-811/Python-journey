#Day 12: Exercise 02: Print Customer + Amount

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

for row in sheet.iter_rows(min_row=2, values_only=True):
    print(f"{row[0]} - {row[1]} €")