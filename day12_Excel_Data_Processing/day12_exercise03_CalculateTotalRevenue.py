#Day 12: Exercise 03: Calculate Total Revenue

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

total_revenue = 0
for row in sheet.iter_rows(min_row=2, values_only=True):
    total_revenue += row[1]

print(f"Total revenue: {total_revenue:.2f} €")