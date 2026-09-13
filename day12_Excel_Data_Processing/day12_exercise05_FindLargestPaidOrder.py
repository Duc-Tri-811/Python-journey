#Day 12: Exercise 05: Find largest paid order

from openpyxl import load_workbook

file_path = r"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day12_Excel_Data_Processing/sales.xlsx"

workbook = load_workbook(file_path)
sheet = workbook["Sales"]

largest_paid_order = 0
for row in sheet.iter_rows(values_only=True):
    if row[2] == "paid":
        if largest_paid_order < row[1]:
            largest_paid_order = row[1]
            largest_paid_customer = row[0]

print(f"Largest paid order: {largest_paid_customer} - {largest_paid_order} €")