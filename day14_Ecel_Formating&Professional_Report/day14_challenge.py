#Day 14: Challenge

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

def generate_professional_sales_report (input_filename, output_filename):

    file_name = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day14_Ecel_Formating&Professional_Report/{input_filename}"

    workbook = load_workbook(file_name)
    sheet = workbook["Sales"]
    report_sheet = workbook.create_sheet("Report")

    sheet["D1"] = "Revenue Category"

    total_orders = 0
    total_revenue = 0
    paid_orders = 0
    paid_order_revenue = 0
    for index,row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        amount = row[1]
        status = row[2]
        

# Calculate KPIs
        total_orders += 1
        total_revenue += amount
        if status == "paid" :
            paid_orders += 1
            paid_order_revenue += amount

#Large/ Small classification
        if amount >= 300:
            category = "Large"
        else:
            category = "Small"

        sheet.cell(row=index, column=4).value = category

    report_sheet.append(["Sales Report"])
    report_sheet["B1"] = "Values"
    report_sheet.append(["Total Orders"])
    report_sheet.append(["Total Revenue"])
    report_sheet.append(["Paid Orders"])
    report_sheet.append(["Paid Order Revenue"])

    report_sheet.cell(row=2, column=2).value = total_orders
    report_sheet.cell(row=3, column=2).value = total_revenue
    report_sheet.cell(row=4, column=2).value = paid_orders
    report_sheet.cell(row=5, column=2).value = paid_order_revenue

# Format Report

    report_sheet["A1"].font = Font(bold=True, size=20)
    report_sheet["A1"].alignment = Alignment(horizontal="center")
    report_sheet.column_dimensions["A"].width = 30
    report_sheet["B1"].font = Font(bold=True, size=20)
    report_sheet["B1"].alignment = Alignment(horizontal="center")
    report_sheet.column_dimensions["B"].width = 30

    for row in report_sheet["A2:B5"]:
        for cell in row:
            cell.font = Font(size=16)

    for row in report_sheet["B2:B5"]:
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

# Format currency
    report_sheet["B3"].number_format = "#,##0.00 €"
    report_sheet["B5"].number_format = "#,##0.00 €"

# Border
    thick = Side(style="thick")
    for row in report_sheet["A1:B1"]:
        for cell in row:
            cell.border = Border(
                top = thick,
                bottom = thick,
                left = thick,
                right = thick
            )

    thin = Side(style="thin")
    for row in report_sheet["A2:B5"]:
        for cell in row:
            cell.border = Border(
                top = thin,
                bottom = thin,
                left = thin,
                right = thin
            )

    new_file_name = rf"C:\Users\Name\OneDrive\Tai lieu\Python\Python-journey\day14_Ecel_Formating&Professional_Report/{output_filename}"

    workbook.save(new_file_name)

generate_professional_sales_report("sales.xlsx", "professional_sales_report.xlsx")